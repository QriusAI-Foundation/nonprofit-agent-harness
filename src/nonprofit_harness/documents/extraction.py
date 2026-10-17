from __future__ import annotations

import io
import time
import zipfile
from dataclasses import dataclass, replace
from typing import TYPE_CHECKING

from nonprofit_harness.core.errors import HarnessError, InvalidRequest
from nonprofit_harness.core.types import Document

if TYPE_CHECKING:
    from nonprofit_harness.config import HarnessConfig

TEXT_TYPES = {"text/plain", "text/markdown", "text/csv", "application/json"}
PDF_TYPES = {"application/pdf"}
DOCX_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

#: UTF-8 spends at most four bytes on one character, so this many bytes always hold
#: at least `max_chars` characters when the file has that many.
_UTF8_MAX_BYTES = 4


class DocumentError(InvalidRequest):
    """A file that cannot be read: corrupt, encrypted, or not what its type says."""

    code = "invalid_document"


class DocumentTooLarge(DocumentError):
    """A file, or what it expands to, is past a configured ceiling."""

    status_code = 413
    code = "document_too_large"


@dataclass(frozen=True, slots=True)
class ExtractionLimits:
    """Ceilings on the work one file may cause.

    The upload byte cap alone does not bound this: a small DOCX can expand many times
    over, and a small PDF can hold thousands of pages. None means no ceiling.
    """

    max_bytes: int | None = 20 * 1024 * 1024
    max_pdf_pages: int | None = 300
    #: python-docx parses every XML part whole, so this, not the character limit,
    #: is what bounds DOCX memory.
    max_docx_expanded_bytes: int | None = 50 * 1024 * 1024
    max_chars: int | None = 2_000_000
    timeout_seconds: float | None = 120.0

    @classmethod
    def from_config(cls, config: HarnessConfig) -> ExtractionLimits:
        return cls(
            max_bytes=config.max_upload_bytes,
            max_pdf_pages=config.max_pdf_pages,
            max_docx_expanded_bytes=config.max_docx_expanded_bytes,
            max_chars=config.max_document_chars,
            timeout_seconds=config.extract_timeout_seconds,
        )

    def check_size(self, size: int, name: str) -> None:
        if self.max_bytes is not None and size > self.max_bytes:
            raise DocumentTooLarge(
                f"{name!r} exceeds the {_megabytes(self.max_bytes)} upload limit"
            )


@dataclass(frozen=True, slots=True)
class Extraction:
    text: str
    pages: int = 1
    truncated: bool = False
    #: Pages actually read, when a page ceiling stopped extraction early.
    pages_read: int | None = None
    #: "pages" or "characters": which ceiling cut the text short.
    truncated_by: str = ""


def supported_media_types() -> set[str]:
    return TEXT_TYPES | PDF_TYPES | DOCX_TYPES


def extract(
    data: bytes,
    *,
    media_type: str,
    name: str = "untitled",
    max_chars: int | None = None,
    limits: ExtractionLimits | None = None,
) -> Document:
    """Turn an uploaded file into the plain text an agent actually receives.

    Agents never see bytes. Keeping extraction here means a new input format is a
    change in one place, not in every agent.

    A file past a ceiling that cannot be read in part (too many bytes, too large once
    unzipped, too slow) raises `DocumentTooLarge`. A file that only has more text or
    pages than allowed comes back cut short, with `truncated` set in its metadata.
    """
    limits = limits or ExtractionLimits(max_bytes=None, max_chars=None)
    if max_chars is not None:
        limits = _with_max_chars(limits, max_chars)
    limits.check_size(len(data), name)

    extraction = _extract_raw(data, media_type=media_type, name=name, limits=limits)
    metadata: dict[str, object] = {
        "pages": extraction.pages,
        "truncated": extraction.truncated,
        "bytes": len(data),
    }
    if extraction.truncated:
        metadata["truncated_by"] = extraction.truncated_by
    if extraction.pages_read is not None:
        metadata["pages_read"] = extraction.pages_read

    return Document(text=extraction.text, name=name, media_type=media_type, metadata=metadata)


def _with_max_chars(limits: ExtractionLimits, max_chars: int) -> ExtractionLimits:
    ceiling = max_chars if limits.max_chars is None else min(max_chars, limits.max_chars)
    return replace(limits, max_chars=ceiling)


def _extract_raw(
    data: bytes, *, media_type: str, name: str, limits: ExtractionLimits
) -> Extraction:
    base_type = media_type.split(";")[0].strip().lower()

    if base_type in TEXT_TYPES:
        return _extract_text(data, limits)
    if base_type in PDF_TYPES:
        return _parse_safely(_extract_pdf, data, name=name, kind="PDF", limits=limits)
    if base_type in DOCX_TYPES:
        return _parse_safely(_extract_docx, data, name=name, kind="DOCX", limits=limits)

    raise InvalidRequest(
        f"Unsupported media type {media_type!r}. Supported: {sorted(supported_media_types())}"
    )


def _parse_safely(parser, data: bytes, *, name: str, kind: str, limits: ExtractionLimits):
    """Run a third-party parser on untrusted bytes and give every failure one shape.

    pypdf and python-docx raise dozens of exception types on malformed input. Without
    this, a corrupt file surfaced as a 500 with a parser traceback.
    """
    try:
        return parser(data, name=name, limits=limits)
    except HarnessError:
        raise
    except Exception as exc:  # noqa: BLE001 - any parser failure means an unreadable file
        raise DocumentError(
            f"Could not read {name!r} as {kind}: the file looks corrupt or is not a "
            f"{kind} ({type(exc).__name__})"
        ) from exc


def _extract_text(data: bytes, limits: ExtractionLimits) -> Extraction:
    if limits.max_chars is None:
        return Extraction(text=data.decode("utf-8", errors="replace"))
    window = data[: limits.max_chars * _UTF8_MAX_BYTES + _UTF8_MAX_BYTES]
    text = window.decode("utf-8", errors="replace")
    if len(text) > limits.max_chars or len(window) < len(data):
        return Extraction(text=text[: limits.max_chars], truncated=True, truncated_by="characters")
    return Extraction(text=text)


def _extract_pdf(data: bytes, *, name: str, limits: ExtractionLimits) -> Extraction:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on optional extra
        raise InvalidRequest(
            "PDF extraction needs the `documents` extra: "
            "pip install 'nonprofit-agent-harness[documents]'"
        ) from exc

    # Started before parsing, so a costly cross-reference table or page tree counts.
    deadline = _Deadline(limits.timeout_seconds, name)
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise DocumentError(f"{name!r} is an encrypted PDF. Remove the password and upload again")

    total = len(reader.pages)
    parts: list[str] = []
    length = 0
    truncated_by = ""
    for index, page in enumerate(reader.pages):
        if limits.max_pdf_pages is not None and index >= limits.max_pdf_pages:
            truncated_by = "pages"
            break
        if limits.max_chars is not None and length >= limits.max_chars:
            truncated_by = "characters"
            break
        deadline.check()
        text = page.extract_text() or ""
        parts.append(text)
        length += len(text) + 2

    text = "\n\n".join(parts).strip()
    if limits.max_chars is not None and len(text) > limits.max_chars:
        text = text[: limits.max_chars]
        truncated_by = truncated_by or "characters"
    return Extraction(
        text=text,
        pages=total,
        truncated=bool(truncated_by),
        pages_read=len(parts) if truncated_by else None,
        truncated_by=truncated_by,
    )


def _extract_docx(data: bytes, *, name: str, limits: ExtractionLimits) -> Extraction:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - depends on optional extra
        raise InvalidRequest(
            "DOCX extraction needs the `documents` extra: "
            "pip install 'nonprofit-agent-harness[documents]'"
        ) from exc

    _check_expansion(data, name=name, ceiling=limits.max_docx_expanded_bytes)
    document = docx.Document(io.BytesIO(data))
    collector = _TextCollector(limits.max_chars)
    for paragraph in document.paragraphs:
        if collector.full or not collector.add(paragraph.text):
            break
    for table in document.tables:
        if collector.full:
            break
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells and not collector.add(" | ".join(cells)):
                break
    return collector.extraction()


def _check_expansion(data: bytes, *, name: str, ceiling: int | None) -> None:
    """Refuse a DOCX whose members would expand past the ceiling, before reading any.

    python-docx loads every part into memory. The zip directory declares each
    member's uncompressed size, and Python's zipfile never returns more than that,
    so summing the declared sizes bounds the expansion without inflating anything.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            expanded = sum(member.file_size for member in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise DocumentError(f"Could not read {name!r} as DOCX: it is not a valid archive") from exc
    if ceiling is not None and expanded > ceiling:
        raise DocumentTooLarge(
            f"{name!r} expands to {_megabytes(expanded)} once unzipped, past the "
            f"{_megabytes(ceiling)} limit"
        )


class _TextCollector:
    """Gathers paragraphs until the character ceiling, then stops taking more."""

    def __init__(self, max_chars: int | None) -> None:
        self.max_chars = max_chars
        self.parts: list[str] = []
        self.length = 0
        self.truncated = False

    @property
    def full(self) -> bool:
        return self.truncated

    def add(self, text: str) -> bool:
        if not text.strip():
            return True
        separator = 1 if self.parts else 0
        if self.max_chars is not None and self.length + separator + len(text) > self.max_chars:
            room = self.max_chars - self.length - separator
            if room > 0:
                self.parts.append(text[:room])
            self.truncated = True
            return False
        self.parts.append(text)
        self.length += separator + len(text)
        return True

    def extraction(self) -> Extraction:
        return Extraction(
            text="\n".join(self.parts).strip(),
            truncated=self.truncated,
            truncated_by="characters" if self.truncated else "",
        )


class _Deadline:
    """A wall-clock ceiling checked between pages.

    It cannot interrupt one slow page, so the page ceiling is what bounds the worst
    case. It stops a long run of merely slow pages from holding a worker for minutes.
    """

    def __init__(self, seconds: float | None, name: str) -> None:
        self.seconds = seconds
        self.name = name
        self.started = time.monotonic()

    def check(self) -> None:
        if self.seconds is None:
            return
        if time.monotonic() - self.started > self.seconds:
            raise DocumentTooLarge(
                f"Reading {self.name!r} took longer than {self.seconds:g} seconds"
            )


def _megabytes(size: int) -> str:
    megabytes = size / (1024 * 1024)
    if megabytes < 1:
        return f"{size:,} bytes"
    return f"{megabytes:,.1f}".removesuffix(".0") + "MB"


__all__ = [
    "DocumentError",
    "DocumentTooLarge",
    "Extraction",
    "ExtractionLimits",
    "extract",
    "supported_media_types",
]
